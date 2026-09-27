from pathlib import Path

from app.ingestion.parsers import parse_document


def test_parse_markdown_tracks_headings(tmp_path: Path):
    md = tmp_path / "doc.md"
    md.write_text(
        "# Leave Policy\n\n"
        "## Carryover Rules\n\n"
        "Employees may carry forward up to 10 days.\n\n"
        "## Sick Leave\n\n"
        "Employees receive 10 paid sick days per year.\n",
        encoding="utf-8",
    )

    blocks = parse_document(md, "md")

    assert len(blocks) == 2
    assert blocks[0].section == "Leave Policy > Carryover Rules"
    assert "carry forward up to 10 days" in blocks[0].text
    assert blocks[1].section == "Leave Policy > Sick Leave"


def test_parse_txt_splits_on_blank_lines(tmp_path: Path):
    txt = tmp_path / "doc.txt"
    txt.write_text("First paragraph.\n\nSecond paragraph with more text.\n", encoding="utf-8")

    blocks = parse_document(txt, "txt")

    assert len(blocks) == 2
    assert blocks[0].text == "First paragraph."
    assert blocks[1].text == "Second paragraph with more text."


def test_parse_docx_detects_headings(tmp_path: Path):
    from docx import Document as DocxDocument

    path = tmp_path / "doc.docx"
    doc = DocxDocument()
    doc.add_heading("Travel Policy", level=0)
    doc.add_paragraph("Economy class is standard for flights under 6 hours.")
    doc.add_heading("Lodging", level=1)
    doc.add_paragraph("Hotel bookings should not exceed $250 per night.")
    doc.save(path)

    blocks = parse_document(path, "docx")

    assert len(blocks) == 2
    assert blocks[0].section == "Travel Policy"
    assert blocks[1].section == "Lodging"
    assert "$250" in blocks[1].text


def test_parse_pdf_extracts_text_with_page_numbers(tmp_path: Path):
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate

    path = tmp_path / "doc.pdf"
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(path), pagesize=LETTER)
    story = [
        Paragraph("This is page one content about vendor onboarding.", styles["BodyText"]),
        PageBreak(),
        Paragraph("This is page two content about vendor termination.", styles["BodyText"]),
    ]
    doc.build(story)

    blocks = parse_document(path, "pdf")

    assert any(b.page_number == 1 and "page one" in b.text for b in blocks)
    assert any(b.page_number == 2 and "page two" in b.text for b in blocks)


def test_unsupported_format_raises(tmp_path: Path):
    from app.ingestion.parsers import UnsupportedFormatError

    path = tmp_path / "doc.xyz"
    path.write_text("data", encoding="utf-8")

    try:
        parse_document(path, "xyz")
        assert False, "expected UnsupportedFormatError"
    except UnsupportedFormatError:
        pass
