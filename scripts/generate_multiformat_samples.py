"""One-off script: converts a few sample_data markdown docs into real
.docx / .pdf / .txt files so ingestion is demonstrably tested against every
supported format, not just Markdown. Run once; the source .md files it
consumes are removed so there's no duplicate content across formats.
"""
import re
import sys
from pathlib import Path

from docx import Document as DocxDocument
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

ROOT = Path(__file__).resolve().parent.parent
SAMPLE_DATA = ROOT / "sample_data"


def parse_md(path: Path) -> list[tuple[str, str]]:
    """Returns a list of (kind, text) where kind is 'h1'|'h2'|'p'."""
    lines = path.read_text(encoding="utf-8").splitlines()
    blocks: list[tuple[str, str]] = []
    buf: list[str] = []

    def flush():
        if buf:
            blocks.append(("p", " ".join(buf).strip()))
            buf.clear()

    for line in lines:
        stripped = line.strip()
        if not stripped:
            flush()
            continue
        if stripped.startswith("## "):
            flush()
            blocks.append(("h2", stripped[3:].strip()))
        elif stripped.startswith("# "):
            flush()
            blocks.append(("h1", stripped[2:].strip()))
        elif stripped.startswith(("- ", "* ")):
            flush()
            blocks.append(("p", stripped[2:].strip()))
        else:
            buf.append(stripped)
    flush()
    return [(k, re.sub(r"[*_`]", "", t)) for k, t in blocks]


def to_docx(md_path: Path, out_path: Path) -> None:
    blocks = parse_md(md_path)
    doc = DocxDocument()
    for kind, text in blocks:
        if kind == "h1":
            doc.add_heading(text, level=0)
        elif kind == "h2":
            doc.add_heading(text, level=1)
        else:
            doc.add_paragraph(text)
    doc.save(out_path)


def to_pdf(md_path: Path, out_path: Path) -> None:
    blocks = parse_md(md_path)
    styles = getSampleStyleSheet()
    doc = SimpleDocTemplate(str(out_path), pagesize=LETTER)
    story = []
    for kind, text in blocks:
        style = {"h1": styles["Title"], "h2": styles["Heading2"]}.get(kind, styles["BodyText"])
        story.append(Paragraph(text, style))
        story.append(Spacer(1, 8))
    doc.build(story)


def to_txt(md_path: Path, out_path: Path) -> None:
    blocks = parse_md(md_path)
    lines = []
    for kind, text in blocks:
        if kind in ("h1", "h2"):
            lines.append(text.upper())
        else:
            lines.append(text)
        lines.append("")
    out_path.write_text("\n".join(lines), encoding="utf-8")


CONVERSIONS = [
    (SAMPLE_DATA / "hr" / "employee_handbook.md", "docx"),
    (SAMPLE_DATA / "hr" / "leave_policy.md", "pdf"),
    (SAMPLE_DATA / "it" / "it_helpdesk_procedures.md", "txt"),
]


def main() -> None:
    for md_path, fmt in CONVERSIONS:
        if not md_path.exists():
            print(f"skip (already converted): {md_path}")
            continue
        out_path = md_path.with_suffix(f".{fmt}")
        if fmt == "docx":
            to_docx(md_path, out_path)
        elif fmt == "pdf":
            to_pdf(md_path, out_path)
        elif fmt == "txt":
            to_txt(md_path, out_path)
        md_path.unlink()
        print(f"converted: {md_path.name} -> {out_path.name}")


if __name__ == "__main__":
    sys.exit(main())
